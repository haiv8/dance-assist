from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
ASSET_DIR = ROOT / "docs" / "ppt-assets"
SRC = ASSET_DIR / "midterm_draft.pptx"
OUT = ROOT / "docs" / "midterm_defense_optimized.pptx"

NAVY = RGBColor(10, 18, 32)
MUTED = RGBColor(99, 116, 139)
TEAL = RGBColor(14, 165, 183)
BLUE_BG = RGBColor(239, 248, 251)
SOFT = RGBColor(246, 250, 252)
CARD = RGBColor(255, 255, 255)
LINE = RGBColor(219, 231, 236)
AMBER = RGBColor(246, 150, 84)


def clear_slide(slide):
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)


def set_background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = SOFT


def add_text(slide, text, x, y, w, h, size=20, bold=False, color=NAVY, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_multiline(slide, lines, x, y, w, h, size=17, color=NAVY, bullet=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        if bullet:
            p.text = f"• {line}"
    return box


def add_round_rect(slide, x, y, w, h, fill=CARD, line=LINE, radius_shape=5):
    shape = slide.shapes.add_shape(radius_shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def add_header(slide, number, title, subtitle):
    add_text(slide, f"{number:02d}", 0.55, 0.35, 0.45, 0.28, size=12, bold=True, color=TEAL)
    add_text(slide, title, 0.55, 0.72, 6.6, 0.45, size=25, bold=True, color=NAVY)
    add_text(slide, subtitle, 0.57, 1.18, 8.4, 0.28, size=10.5, color=MUTED)
    add_text(slide, "中期答辩 · 基于视频分析的舞蹈学习辅助平台", 8.6, 0.38, 4.0, 0.24, size=9.5, color=MUTED, align=PP_ALIGN.RIGHT)


def add_image_fit(slide, image_path, x, y, w, h):
    with Image.open(image_path) as img:
        iw, ih = img.size
    box_w = Inches(w)
    box_h = Inches(h)
    scale = min(box_w / iw, box_h / ih)
    pic_w = int(iw * scale)
    pic_h = int(ih * scale)
    left = int(Inches(x) + (box_w - pic_w) / 2)
    top = int(Inches(y) + (box_h - pic_h) / 2)
    return slide.shapes.add_picture(str(image_path), left, top, width=pic_w, height=pic_h)


def add_screenshot_card(slide, image_name, x, y, w, h, caption, tag):
    add_round_rect(slide, x, y, w, h, fill=CARD, line=LINE)
    add_image_fit(slide, ASSET_DIR / image_name, x + 0.12, y + 0.16, w - 0.24, h - 0.68)
    add_text(slide, tag, x + 0.22, y + h - 0.43, 1.5, 0.18, size=8.5, bold=True, color=TEAL)
    add_text(slide, caption, x + 0.82, y + h - 0.46, w - 1.05, 0.24, size=10.5, bold=True, color=NAVY)


def rebuild_slide_9(slide):
    clear_slide(slide)
    set_background(slide)
    add_header(
        slide,
        9,
        "前端闭环 1：素材准备到发起分析",
        "把上传、整理、配对和发起分析收敛在同一条任务链路中，减少演示时的页面跳转成本。",
    )
    add_screenshot_card(slide, "library.png", 0.55, 1.72, 6.0, 4.65, "素材库：教师/学员素材、推荐配对、快速带入分析", "STEP 01")
    add_screenshot_card(slide, "analysis.png", 6.78, 1.72, 6.0, 4.65, "开始分析：流程状态、当前组合、对照舞台入口", "STEP 02")
    add_text(slide, "答辩讲法：先说明用户不需要理解算法细节，只需要准备教师与学员视频，系统会组织成一次可追踪的分析任务。", 0.65, 6.65, 12.1, 0.32, size=12, color=MUTED)


def rebuild_slide_10(slide):
    clear_slide(slide)
    set_background(slide)
    add_header(
        slide,
        10,
        "前端闭环 2：记录复盘与问题回看",
        "分析完成后沉淀为记录、问题片段和对照舞台，形成“知道哪里错、回看哪里、怎么练”的闭环。",
    )
    add_screenshot_card(slide, "records.png", 0.55, 1.62, 6.1, 4.86, "分析记录：按状态筛选、批量删除、查看得分与问题片段", "RECORDS")
    add_screenshot_card(slide, "compare_stage.png", 6.78, 1.62, 6.0, 4.86, "对照舞台：教师示范与学员练习并排回放，支持片段复盘", "REPLAY")
    add_text(slide, "产品重点：不是只给一个分数，而是把算法输出转化为可执行的训练反馈。", 0.65, 6.68, 12.1, 0.32, size=12, bold=True, color=TEAL)


def rebuild_slide_11(slide):
    clear_slide(slide)
    set_background(slide)
    add_header(
        slide,
        11,
        "工程稳定性：面向答辩演示的可维护改造",
        "近期优化从“功能能跑”推进到“演示稳定、数据可解释、后续可维护”。",
    )
    add_screenshot_card(slide, "settings.png", 0.55, 1.62, 6.2, 4.95, "系统设置：本地服务、存储、模型文件、Redis/PostgreSQL 状态检查", "STATUS")

    add_round_rect(slide, 7.08, 1.62, 5.65, 4.95, fill=CARD, line=LINE)
    add_text(slide, "当前工程优化重点", 7.42, 1.95, 3.8, 0.34, size=18, bold=True, color=NAVY)
    bullets = [
        "记录中心默认限制最近 50 条，避免历史数据过多拖慢列表。",
        "分析完成后生成 issues.json，记录页优先读取轻量问题索引。",
        "旧记录没有索引时自动 fallback 解析 report，并顺手补齐索引。",
        "AI 助教只负责解释结构化报告，不参与动作误差判定。",
        "桌面端启动、系统健康检查和本地服务状态已面向演示场景收敛。",
    ]
    add_multiline(slide, bullets, 7.42, 2.58, 4.85, 2.95, size=13.5, color=NAVY, bullet=True)
    add_round_rect(slide, 7.42, 5.72, 4.8, 0.52, fill=BLUE_BG, line=RGBColor(186, 226, 235))
    add_text(slide, "答辩边界：核心判断来自姿态估计 + DTW 对齐 + 指标计算，AI 是辅助解读层。", 7.65, 5.88, 4.35, 0.18, size=9.5, bold=True, color=TEAL)


def main():
    prs = Presentation(str(SRC))
    rebuild_slide_9(prs.slides[8])
    rebuild_slide_10(prs.slides[9])
    rebuild_slide_11(prs.slides[10])
    prs.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
